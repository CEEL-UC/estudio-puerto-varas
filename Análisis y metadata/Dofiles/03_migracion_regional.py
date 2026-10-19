# Migración estudiantil con óptica regional (Región de Los Lagos), 2015-2025.
# Lee los archivos NACIONALES de Matrícula por estudiante (Mineduc) para saber dónde estaba cada alumno el año
# anterior y el siguiente en cualquier comuna de Chile. Así un alumno que llega a Puerto Varas desde Llanquihue
# cuenta como entrante de otra comuna de la región, y no como alguien "de fuera" (como pasaba al mirar solo 4 comunas).
#
# Definiciones (alumno del tramo regular, 1° básico a 4° medio, matriculado en la comuna c el año t):
#   Ingreso 1° básico: cursa 1° básico en t (comienza el tramo regular).
#   Se mantiene: en t-1 estaba matriculado (en cualquier nivel) en un colegio de la misma comuna c.
#   Entrante de otra comuna de Los Lagos: en t-1 estaba en un colegio de otra comuna de la región.
#   Entrante de otra región: en t-1 estaba en un colegio de una comuna de otra región.
#   Sin registro el año anterior: no aparece matriculado en ningún colegio de Chile en t-1 (llega del extranjero,
#     vuelve al sistema, o venía de un jardín que no informa a Mineduc).
#   Salidas (solo hasta 4° medio exclusive, t <= 2024): no está en la comuna c en t+1; se clasifican en
#     a otra comuna de Los Lagos / a otra región / sin registro el año siguiente.
# Escribe un JSON con todas las series para el sitio. Correr con:
#   C:\Users\gonza\AppData\Local\Programs\Python\Python312\python.exe 03_migracion_regional.py
import glob, json, os
import polars as pl
from openpyxl import load_workbook

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
ORIGEN = os.path.join(BASE, 'Datos', 'Mineduc', 'Matricula por estudiante')
CACHE = os.path.join(BASE, 'Datos', '[merge] Panel Puerto Varas', 'nacional_liviano')
DPA = os.path.join(BASE, 'Datos', 'Censo 2024', 'documentacion_manzana', 'diccionario_variables_glosas_censo2024.xlsx')
SALIDA = os.path.join(BASE, 'Análisis y metadata', 'Outputs', 'migracion_regional_sitio.json')
ANIOS = list(range(2014, 2026))
REPORTE = list(range(2015, 2026))
REGION = 10
ESTUDIO = {10109: 'Puerto Varas', 10101: 'Puerto Montt', 10105: 'Frutillar', 10301: 'Osorno'}
DEPE = {1: 'Municipal', 2: 'Part. Subvencionado', 3: 'Part. Pagado', 4: 'Corp. Adm. Delegada', 5: 'Serv. Local Educ.'}
c = pl.col
os.makedirs(CACHE, exist_ok=True)


def liviano(anio):
    # versión liviana del archivo nacional del año: una fila por matrícula, solo las columnas necesarias
    dst = os.path.join(CACHE, f'mat_{anio}.parquet')
    if os.path.exists(dst):
        return pl.read_parquet(dst, glob=False)  # la carpeta tiene corchetes: sin glob
    f = [x for x in glob.glob(os.path.join(ORIGEN, str(anio), '*')) if x.lower().endswith('.csv')][0]
    # algunos años traen los nombres de columna en minúscula y/o con BOM al inicio
    hdr = [x.strip().upper() for x in open(f, encoding='latin-1').readline().strip().split(';')]
    hdr[0] = hdr[0][-4:] if hdr[0].endswith('AGNO') else hdr[0]
    cols = ['RBD', 'NOM_RBD', 'COD_REG_RBD', 'COD_COM_RBD', 'COD_DEPE2', 'COD_ENSE', 'COD_GRADO', 'MRUN']
    d = pl.read_csv(f, separator=';', encoding='utf8-lossy', infer_schema=False, quote_char=None,
                    columns=[hdr.index(x) for x in cols], new_columns=cols)
    d = d.with_columns([c(x).str.strip_chars().cast(pl.Int64, strict=False) for x in cols if x != 'NOM_RBD'])
    d = d.with_columns(curso=pl.when(c('COD_ENSE') == 110).then(c('COD_GRADO'))
                       .when(c('COD_ENSE').is_in([310, 410, 510, 610, 710, 810])).then(c('COD_GRADO') + 8))
    d = d.with_columns(regular=c('curso').is_between(1, 12)).with_columns(anio=pl.lit(anio))
    d.write_parquet(dst)
    return d


datos = {}
for a in ANIOS:
    datos[a] = liviano(a)
    x = datos[a]
    print(a, 'filas', x.height, '| MRUN vacío', x['MRUN'].null_count(), '| tramo regular', x['regular'].sum())


def ubicacion(a):
    # una comuna por alumno en el año a: si tiene varias matrículas se prefiere la del tramo regular
    return (datos[a].filter(c('MRUN').is_not_null()).sort(['MRUN', 'regular'], descending=[False, True])
            .unique('MRUN', keep='first').select('MRUN', c('COD_COM_RBD').alias('com')))


# alumnos del tramo regular en la región, con su comuna el año anterior y el siguiente
filas = []
for t in REPORTE:
    reg = (datos[t].filter((c('COD_REG_RBD') == REGION) & c('regular') & c('MRUN').is_not_null())
           .sort(['MRUN', 'RBD']).unique('MRUN', keep='first'))
    reg = reg.join(ubicacion(t - 1).rename({'com': 'com_ant'}), on='MRUN', how='left')
    if t + 1 in datos:
        reg = reg.join(ubicacion(t + 1).rename({'com': 'com_sig'}), on='MRUN', how='left')
    else:
        reg = reg.with_columns(com_sig=pl.lit(None, pl.Int64))
    sin_mrun = datos[t].filter((c('COD_REG_RBD') == REGION) & c('regular') & c('MRUN').is_null()).height
    print(t, 'alumnos tramo regular región', reg.height, '| sin MRUN (excluidos)', sin_mrun)
    filas.append(reg)
al = pl.concat(filas, how='diagonal_relaxed')

al = al.with_columns(
    origen=pl.when(c('curso') == 1).then(pl.lit('Ingreso 1° básico'))
    .when(c('com_ant').is_null()).then(pl.lit('Sin registro el año anterior'))
    .when(c('com_ant') == c('COD_COM_RBD')).then(pl.lit('Se mantiene'))
    .when(c('com_ant') // 1000 == REGION).then(pl.lit('Entrante de otra comuna de Los Lagos'))
    .otherwise(pl.lit('Entrante de otra región')),
    destino=pl.when((c('anio') == max(REPORTE)) | (c('curso') == 12)).then(None)
    .when(c('com_sig').is_null()).then(pl.lit('Sin registro el año siguiente'))
    .when(c('com_sig') == c('COD_COM_RBD')).then(pl.lit('Sigue en la comuna'))
    .when(c('com_sig') // 1000 == REGION).then(pl.lit('Salida a otra comuna de Los Lagos'))
    .otherwise(pl.lit('Salida a otra región')),
    reg_ant=pl.when(c('curso') == 1).then(pl.lit('Ingreso 1° básico'))
    .when(c('com_ant').is_null()).then(pl.lit('Sin registro el año anterior'))
    .when(c('com_ant') // 1000 == REGION).then(pl.lit('Se mantiene en la región'))
    .otherwise(pl.lit('Entrante de otra región')),
    reg_sig=pl.when((c('anio') == max(REPORTE)) | (c('curso') == 12)).then(None)
    .when(c('com_sig').is_null()).then(pl.lit('Sin registro el año siguiente'))
    .when(c('com_sig') // 1000 == REGION).then(pl.lit('Sigue en la región'))
    .otherwise(pl.lit('Salida a otra región')),
    depe=c('COD_DEPE2').replace_strict(DEPE, default=None, return_dtype=pl.Utf8),
)

# nombres de comunas con tildes, desde la división político administrativa del Censo 2024
wb = load_workbook(DPA, read_only=True)
NOM = {}
for r in list(wb['DPA'].iter_rows(values_only=True))[1:]:
    if r[4]:
        nombre = ' '.join(p if p in ('de', 'del', 'la', 'las', 'los', 'y') else p.capitalize() for p in str(r[5]).strip().lower().split())
        NOM[int(r[4])] = nombre[0].upper() + nombre[1:]
NOM.update(ESTUDIO)

ORIG = ['Se mantiene', 'Entrante de otra comuna de Los Lagos', 'Entrante de otra región', 'Sin registro el año anterior', 'Ingreso 1° básico']
DEST = ['Salida a otra comuna de Los Lagos', 'Salida a otra región', 'Sin registro el año siguiente']
ORIG_REG = ['Se mantiene en la región', 'Entrante de otra región', 'Sin registro el año anterior', 'Ingreso 1° básico']
DEST_REG = ['Salida a otra región', 'Sin registro el año siguiente']


def serie(df, col, cats, anios_validos=None):
    t = df.group_by(['anio', col]).len()
    out = {}
    for k in cats:
        v = {r[0]: r[2] for r in t.filter(c(col) == k).rows()}
        out[k] = [(v.get(a, 0) if (anios_validos is None or a in anios_validos) else None) for a in REPORTE]
    return out


VAL_SAL = set(REPORTE[:-1])
comunas = sorted(al['COD_COM_RBD'].unique().to_list(), key=lambda x: (x != 10109, NOM.get(x, str(x))))
res = {'anios': REPORTE, 'comunas': [], 'flujo': {}, 'flujoDepe': {}}
for cod in comunas:
    nom = NOM.get(cod, str(cod))
    sub = al.filter(c('COD_COM_RBD') == cod)
    res['comunas'].append({'cod': cod, 'n': nom, 'estudio': cod in ESTUDIO})
    f = serie(sub, 'origen', ORIG)
    f.update(serie(sub, 'destino', DEST, VAL_SAL))
    res['flujo'][nom] = f
    res['flujoDepe'][nom] = {}
    for dp in ['Municipal', 'Part. Subvencionado', 'Part. Pagado', 'Serv. Local Educ.', 'Corp. Adm. Delegada']:
        sd = sub.filter(c('depe') == dp)
        g = serie(sd, 'origen', ORIG)
        g.update(serie(sd, 'destino', DEST, VAL_SAL))
        res['flujoDepe'][nom][dp] = g
reg = serie(al, 'reg_ant', ORIG_REG)
reg.update(serie(al, 'reg_sig', DEST_REG, VAL_SAL))
res['region'] = reg

# por colegio (solo los de las 4 comunas del mapa): entradas, ingreso 1° básico, se mantiene, salidas, y top 5
# comunas de origen y destino
est = al.filter(c('COD_COM_RBD').is_in(list(ESTUDIO)))
res['colegio'] = {}
for r in est.group_by(['RBD', 'anio']).agg(
        e=c('origen').is_in(ORIG[1:4]).sum(), i=(c('origen') == 'Ingreso 1° básico').sum(),
        k=(c('origen') == 'Se mantiene').sum(), s=c('destino').is_in(DEST).sum()).rows(named=True):
    res['colegio'].setdefault(str(r['RBD']), {})[str(r['anio'])] = {
        'e': r['e'], 'i': r['i'], 'k': r['k'], 's': (r['s'] if r['anio'] in VAL_SAL else None)}


def top5(df, col_com, filtro):
    out = {}
    t = df.filter(filtro).group_by(['RBD', 'anio', col_com]).len().sort(['RBD', 'anio', 'len'], descending=[False, False, True])
    for (rbd, anio), g in t.group_by(['RBD', 'anio'], maintain_order=True):
        out.setdefault(str(rbd), {})[str(anio)] = [{'c': NOM.get(x, str(x)), 'n': n} for x, n in g.select(col_com, 'len').rows()[:5]]
    return out


res['origenTop5'] = top5(est, 'com_ant', c('origen').is_in(ORIG[1:3]))
res['destinoTop5'] = top5(est, 'com_sig', c('destino').is_in(DEST[:2]))
res['nombres_rbd'] = {str(r[0]): r[1] for r in est.unique('RBD').select('RBD', 'NOM_RBD').rows()}

# comunas de origen de los entrantes, por comuna de llegada, año y dependencia del colegio de llegada
# (M = Municipal o SLEP, S = Part. Subvencionado, P = Part. Pagado, C = Corp. Adm. Delegada).
# i = desde otra comuna de Los Lagos, x = desde otra región. Formato compacto: [código de comuna, alumnos].
DEPK = {'Municipal': 'M', 'Serv. Local Educ.': 'M', 'Part. Subvencionado': 'S', 'Part. Pagado': 'P', 'Corp. Adm. Delegada': 'C'}
ent = al.filter(c('origen').is_in(ORIG[1:3])).with_columns(
    dk=c('depe').replace_strict(DEPK, default='C', return_dtype=pl.Utf8),
    tipo=pl.when(c('origen') == ORIG[2]).then(pl.lit('x')).otherwise(pl.lit('i')))
oc = {}
for (cod, anio, dk, tipo), g in ent.group_by(['COD_COM_RBD', 'anio', 'dk', 'tipo']):
    lista = g.group_by('com_ant').len().sort('len', descending=True).rows()
    oc.setdefault(NOM.get(cod, str(cod)), {}).setdefault(str(anio), {}).setdefault(dk, {})[tipo] = [[x, n] for x, n in lista]
res['origenComunas'] = oc
# nombre y región de cada comuna de origen; los códigos antiguos de Ñuble (8401-8421, antes de 2018) no están en
# la DPA 2024 y se rotulan con la región que tenían entonces
REGNOM = {1: 'Tarapacá', 2: 'Antofagasta', 3: 'Atacama', 4: 'Coquimbo', 5: 'Valparaíso', 6: "O'Higgins", 7: 'Maule',
          8: 'Biobío', 9: 'La Araucanía', 10: 'Los Lagos', 11: 'Aysén', 12: 'Magallanes', 13: 'Metropolitana',
          14: 'Los Ríos', 15: 'Arica y Parinacota', 16: 'Ñuble'}
cods = {x for lst in [v2 for v in oc.values() for a in v.values() for d in a.values() for v2 in d.values()] for x, _ in lst}
sin_nombre = sorted(x for x in cods if x not in NOM)
if sin_nombre:
    # el nombre se toma del propio archivo de Mineduc de 2017 (último año con los códigos antiguos de Ñuble)
    f17 = [x for x in glob.glob(os.path.join(ORIGEN, '2017', '*')) if x.lower().endswith('.csv')][0]
    hdr = [x.strip().upper() for x in open(f17, encoding='latin-1').readline().strip().split(';')]
    n17 = pl.read_csv(f17, separator=';', encoding='utf8-lossy', infer_schema=False, quote_char=None,
                      columns=[hdr.index('COD_COM_RBD'), hdr.index('NOM_COM_RBD')], new_columns=['cod', 'nom']).unique()
    import unicodedata
    plano = lambda t: unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode().lower().strip()
    nuble = {plano(v): (k, v) for k, v in NOM.items() if k // 1000 == 16}
    REG_ANT = {}
    for cod, nom in n17.rows():
        cod = int(cod)
        if cod in sin_nombre:
            # se usa el nombre actual (con tildes) de la misma comuna en la Región de Ñuble, si calza
            nuevo = nuble.get(plano(nom))
            NOM[cod] = nuevo[1] if nuevo else nom.strip().title()
            REG_ANT[cod] = 'Ñuble' if nuevo else REGNOM.get(cod // 1000)
res['comunasOrigen'] = {str(x): [NOM.get(x, f'Comuna {x}'), (REG_ANT.get(x) if sin_nombre and x in REG_ANT else REGNOM.get(x // 1000, f'Región {x // 1000}'))] for x in sorted(cods)}
print('comunas de origen:', len(cods), '| sin nombre:', [x for x in cods if x not in NOM], '| Ñuble antiguas:', {x: NOM.get(x) for x in sin_nombre})

json.dump(res, open(SALIDA, 'w', encoding='utf8'), ensure_ascii=False)
print('comunas:', len(comunas), '| JSON:', SALIDA)
pv = res['flujo']['Puerto Varas']
for k, v in pv.items():
    print('  PV', k, v)
