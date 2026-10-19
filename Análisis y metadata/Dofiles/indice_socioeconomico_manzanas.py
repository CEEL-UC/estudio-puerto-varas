# Índice socioeconómico (proxy) por manzana, Censo 2024: Excel con fórmulas vivas, una fila por polígono.
# Usa exactamente la misma definición que prep_manzanas_censo.py (la que alimenta el mapa del sitio).
# Correr con: C:\Users\gonza\AppData\Local\Programs\Python\Python312\python.exe indice_socioeconomico_manzanas.py
import glob, json, os
import polars as pl
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), '..', '..')
CARTO = os.path.join(BASE, 'Datos', 'Censo 2024', 'cartografia')
SALIDA_DIR = os.path.join(BASE, 'Análisis y metadata', 'Outputs', 'Indice socioeconomico manzanas')
SALIDA = os.environ.get('SALIDA_XLSX') or os.path.join(SALIDA_DIR, 'indice_socioeconomico_manzanas_censo2024_formulas.xlsx')
DATOS_SITIO = r'C:\Users\gonza\repos\estudio-puerto-varas-informe\datos_manzanas.js'

COMUNAS = {10109: 'Puerto Varas', 10101: 'Puerto Montt', 10105: 'Frutillar', 10301: 'Osorno'}
c = pl.col


def s(*cols):
    return pl.sum_horizontal([c(x).fill_null(0) for x in cols])


def tasa(num, den):
    return pl.when(num.is_not_null() & den.is_not_null() & (den > 0)).then((num / den * 100).round(1, mode='half_away_from_zero')).otherwise(None)


def suma_o_null(*cols):
    return pl.when(pl.any_horizontal([c(x).is_not_null() for x in cols])).then(s(*cols)).otherwise(None)


TIPO_VIV = ['n_tipo_viv_casa', 'n_tipo_viv_depto', 'n_tipo_viv_indigena', 'n_tipo_viv_pieza', 'n_tipo_viv_mediagua',
            'n_tipo_viv_movil', 'n_tipo_viv_otro']
PAREDES = ['n_mat_paredes_hormigon', 'n_mat_paredes_albanileria', 'n_mat_paredes_tabique_forrado',
           'n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios']
mayores18 = c('n_per') - s('n_edad_0_5', 'n_edad_6_13', 'n_edad_14_17')
viv_ocup = s(*TIPO_VIV)

# componente: (expresión, signo, nombre para el Excel)
COMP = {
    'esc': (c('prom_escolaridad18').round(1, mode='half_away_from_zero'), 1, 'Escolaridad promedio 18+ (años)'),
    'sup': (tasa(c('n_cine_terciaria_maestria_doctorado'), mayores18), 1, '% educación superior (18+)'),
    'prof': (tasa(suma_o_null('n_ciuo_1', 'n_ciuo_2'), c('n_ocupado')), 1, '% directivos y profesionales (ocupados)'),
    'compu': (tasa(c('n_serv_compu'), c('n_hog')), 1, '% hogares con computador'),
    'hacin': (tasa(c('n_viv_hacinadas'), viv_ocup), -1, '% viviendas hacinadas'),
    'precar': (tasa(suma_o_null('n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios'), s(*PAREDES)), -1, '% viviendas con paredes precarias'),
    'alleg': (tasa(c('n_hog_allegados'), viv_ocup), -1, 'Hogares allegados por 100 viviendas'),
}
MIN_COMP = 4

# variables INE que entran como insumo al Excel (el resto se calcula con fórmulas dentro del libro)
INSUMOS = ['prom_escolaridad18', 'n_per', 'n_edad_0_5', 'n_edad_6_13', 'n_edad_14_17', 'n_cine_terciaria_maestria_doctorado',
           'n_ocupado', 'n_ciuo_1', 'n_ciuo_2', 'n_hog', 'n_serv_compu', 'n_viv_hacinadas', 'n_hog_allegados'] + TIPO_VIV + PAREDES

partes = []
for capa in ['Manzanas', 'Entidades']:
    f = glob.glob(os.path.join(CARTO, f'*_{capa}.parquet'))[0]
    d = pl.read_parquet(f).drop('SHAPE').with_columns(c('CUT').cast(pl.Int64))
    d = d.filter(c('CUT').is_in(list(COMUNAS)) & (c('MZ_BASE_CENSO') == 1))
    num = [x for x in d.columns if x.startswith(('n_', 'prom_'))]
    d = d.with_columns([c(x).cast(pl.Float64) for x in num])
    d = d.with_columns(**{k: v[0] for k, v in COMP.items()})
    d = d.with_columns(tipo=pl.lit('Manzana urbana' if capa == 'Manzanas' else 'Entidad rural'))
    partes.append(d.select(['MANZENT', 'CUT', 'tipo', 'LOCALIDAD', 'ENTIDAD', 'n_vp'] + INSUMOS + list(COMP)))
d = pl.concat(partes).with_columns(comuna=c('CUT').replace_strict(COMUNAS, return_dtype=pl.Utf8))

# parámetros de estandarización: media y desviación estándar muestral sobre los 6.369 polígonos (4 comunas)
par = {k: {'media': d[k].mean(), 'de': d[k].std(), 'n': d[k].drop_nulls().len(), 'signo': v[1], 'nombre': v[2]} for k, v in COMP.items()}
d = d.with_columns([((c(k) - par[k]['media']) / par[k]['de'] * par[k]['signo']).alias('z_' + k) for k in COMP])
zc = ['z_' + k for k in COMP]
d = d.with_columns(n_comp=pl.sum_horizontal([c(x).is_not_null() for x in zc]).cast(pl.Int8))
# el índice se promedia con los z sin redondear (igual que el sitio); los z se redondean solo para mostrarlos
d = d.with_columns(indice=pl.when(c('n_comp') >= MIN_COMP).then(pl.mean_horizontal(zc).round(2, mode='half_away_from_zero')).otherwise(None))
d = d.with_columns([c(x).round(3, mode='half_away_from_zero') for x in zc])

# percentil sobre las 4 comunas: % de polígonos con índice estrictamente menor
valid = d['indice'].drop_nulls().to_list()
srt = sorted(valid)
import bisect
from decimal import Decimal, ROUND_HALF_UP
redondea1 = lambda v: float(Decimal(repr(v)).quantize(Decimal('0.1'), ROUND_HALF_UP))  # mismo redondeo que ROUND de Excel
d = d.with_columns(percentil=c('indice').map_elements(lambda x: redondea1(100 * bisect.bisect_left(srt, x) / len(srt)), return_dtype=pl.Float64))


def quintiles_mapa(vals):
    # mismo algoritmo que el mapa del sitio: cortes en vals[floor(i*n/5)], sin repetir ni igualar al mínimo
    vals = sorted(vals)
    q = []
    for i in range(1, 5):
        x = vals[(i * len(vals)) // 5]
        if x > vals[0] and x not in q:
            q.append(x)
    return q


def asignar(x, cortes):
    if x is None:
        return None
    return 1 + sum(1 for b in cortes if x >= b)


filas_q = []
cortes_info = []
for com in COMUNAS.values():
    for alcance, filtro in [('todos', None), ('urbano', 'Manzana urbana')]:
        sub = d.filter((c('comuna') == com) & ((c('tipo') == filtro) if filtro else pl.lit(True)))
        cortes = quintiles_mapa(sub['indice'].drop_nulls().to_list())
        cortes_info.append({'comuna': com, 'alcance': alcance, 'cortes': cortes})
        for mz, x in zip(sub['MANZENT'].to_list(), sub['indice'].to_list()):
            filas_q.append({'MANZENT': mz, 'alcance': alcance, 'q': asignar(x, cortes)})
q = pl.DataFrame(filas_q).pivot(on='alcance', index='MANZENT', values='q')
d = d.join(q.rename({'todos': 'quintil_comuna', 'urbano': 'quintil_comuna_urbano'}), on='MANZENT', how='left')

# validación contra los datos publicados en el sitio
txt = open(DATOS_SITIO, encoding='utf8').read()
mzk = json.loads(txt[txt.index('var MZK=') + 8: txt.index(';')])
mz = json.loads(txt[txt.index('var MZ=') + 7: txt.rindex(';')])
pos = 2 + mzk.index('ise')
sitio = {int(r['id']): r['v'][pos] for r in mz}
dif = sum(1 for a, b in zip(d['MANZENT'].to_list(), d['indice'].to_list()) if sitio.get(int(a)) != b)
print('polígonos:', d.height, '| con índice:', len(valid), '| diferencias contra el sitio:', dif)

# ---- Excel con fórmulas vivas
# Los datos del INE van como valores (insumos). Componentes, puntajes z, parámetros, índice, percentil, cortes,
# quintiles y resumen son fórmulas. openpyxl escribe las funciones en inglés; Excel las muestra en el idioma
# del usuario (PROMEDIO, DESVEST.M, etc.). Después se abre el libro con Excel para que calcule y quede guardado
# con los valores, y se comparan contra el cálculo en Python.
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Border, Side
from openpyxl.worksheet.table import Table

d = d.sort(['comuna', 'tipo', 'MANZENT'])
N = d.height
R0, R1 = 2, d.height + 1  # filas de datos en la hoja Manzanas

ETIQ_INS = {
    'prom_escolaridad18': 'prom_escolaridad18 (años)', 'n_per': 'n_per (personas)', 'n_edad_0_5': 'n_edad_0_5',
    'n_edad_6_13': 'n_edad_6_13', 'n_edad_14_17': 'n_edad_14_17', 'n_cine_terciaria_maestria_doctorado': 'n_cine_terciaria_maestria_doctorado',
    'n_ocupado': 'n_ocupado', 'n_ciuo_1': 'n_ciuo_1', 'n_ciuo_2': 'n_ciuo_2', 'n_hog': 'n_hog (hogares)', 'n_serv_compu': 'n_serv_compu',
    'n_viv_hacinadas': 'n_viv_hacinadas', 'n_hog_allegados': 'n_hog_allegados',
}
cols = ['MANZENT', 'Comuna', 'Tipo de polígono', 'Localidad', 'Entidad', 'n_vp (viviendas particulares)'] \
    + [ETIQ_INS.get(x, x) for x in INSUMOS] \
    + [v[2] for v in COMP.values()] + ['z ' + v[2] for v in COMP.values()] \
    + ['Componentes con dato', 'Índice socioeconómico', 'Percentil (4 comunas)',
       'Quintil en la comuna (todos los polígonos)', 'Quintil en la comuna (solo manzanas urbanas)']
L = {}  # nombre lógico -> letra de columna
for i, nombre in enumerate(['MANZENT', 'comuna', 'tipo', 'loc', 'ent', 'n_vp'] + INSUMOS + list(COMP) + ['z_' + k for k in COMP]
                           + ['n_comp', 'indice', 'percentil', 'q_todos', 'q_urb'], start=1):
    L[nombre] = get_column_letter(i)


def rng(k):
    return f"Manzanas!${L[k]}${R0}:${L[k]}${R1}"


def ref(k, r):
    return f"{L[k]}{r}"


PAR = "'Parámetros'"
ORDEN_COMP = list(COMP)


def formulas_fila(r):
    f = {}
    x = lambda k: ref(k, r)
    f['esc'] = f'=IF({x("prom_escolaridad18")}="","",ROUND({x("prom_escolaridad18")},1))'
    den18 = f'({x("n_per")}-({x("n_edad_0_5")}+{x("n_edad_6_13")}+{x("n_edad_14_17")}))'
    f['sup'] = f'=IF(OR({x("n_cine_terciaria_maestria_doctorado")}="",{den18}<=0),"",ROUND(100*{x("n_cine_terciaria_maestria_doctorado")}/{den18},1))'
    f['prof'] = f'=IF(OR(AND({x("n_ciuo_1")}="",{x("n_ciuo_2")}=""),N({x("n_ocupado")})<=0),"",ROUND(100*({x("n_ciuo_1")}+{x("n_ciuo_2")})/{x("n_ocupado")},1))'
    f['compu'] = f'=IF(OR({x("n_serv_compu")}="",N({x("n_hog")})<=0),"",ROUND(100*{x("n_serv_compu")}/{x("n_hog")},1))'
    viv = f'SUM({ref(TIPO_VIV[0], r)}:{ref(TIPO_VIV[-1], r)})'
    f['hacin'] = f'=IF(OR({x("n_viv_hacinadas")}="",{viv}<=0),"",ROUND(100*{x("n_viv_hacinadas")}/{viv},1))'
    par_ = f'SUM({ref(PAREDES[0], r)}:{ref(PAREDES[-1], r)})'
    a, b, cc = (ref(k, r) for k in ['n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios'])
    f['precar'] = f'=IF(OR(AND({a}="",{b}="",{cc}=""),{par_}<=0),"",ROUND(100*({a}+{b}+{cc})/{par_},1))'
    f['alleg'] = f'=IF(OR({x("n_hog_allegados")}="",{viv}<=0),"",ROUND(100*{x("n_hog_allegados")}/{viv},1))'
    for i, k in enumerate(ORDEN_COMP):
        pr = 2 + i  # fila del componente en la hoja Parámetros
        f['z_' + k] = f'=IF({x(k)}="","",{PAR}!$C${pr}*({x(k)}-{PAR}!$D${pr})/{PAR}!$E${pr})'
    zr = f'{ref("z_" + ORDEN_COMP[0], r)}:{ref("z_" + ORDEN_COMP[-1], r)}'
    f['n_comp'] = f'=COUNT({zr})'
    f['indice'] = f'=IF({x("n_comp")}>={PAR}!$C$11,ROUND(AVERAGE({zr}),2),"")'
    f['percentil'] = f'=IF({x("indice")}="","",ROUND(100*COUNTIF({rng("indice")},"<"&{x("indice")})/COUNT({rng("indice")}),1))'
    # quintil: 1 + número de cortes válidos de su comuna menores o iguales al índice (hoja Cortes de quintiles)
    for key, alc in [('q_todos', 'Todos los polígonos'), ('q_urb', 'Solo manzanas urbanas')]:
        cond_urb = '' if key == 'q_todos' else f'{x("tipo")}<>"Manzana urbana",'
        f[key] = (f"=IF(OR({cond_urb}{x('indice')}=\"\"),\"\",1+SUMPRODUCT(('Cortes de quintiles'!$A$2:$A$9={x('comuna')})"
                  f"*('Cortes de quintiles'!$B$2:$B$9=\"{alc}\")*('Cortes de quintiles'!$E$2:$H$9<={x('indice')})"
                  f"*('Cortes de quintiles'!$I$2:$L$9=TRUE)))")
    return f


wb = Workbook()
ws = wb.active
ws.title = 'Manzanas'
ws.append(cols)
ins_vals = d.select(['MANZENT', 'comuna', 'tipo', 'LOCALIDAD', 'ENTIDAD', 'n_vp'] + INSUMOS).rows()
for i, fila in enumerate(ins_vals):
    r = R0 + i
    mz, com, tipo, loc, ent, nvp, *ins = fila
    vals = [str(int(mz)), com, tipo, loc, ent, None if nvp is None else int(nvp)]
    vals += [None if v is None else (v if k == 'prom_escolaridad18' else int(v)) for k, v in zip(INSUMOS, ins)]
    ws.append(vals)
    for k, frm in formulas_fila(r).items():
        ws[ref(k, r)] = frm

# formato: encabezados por bloque, anchos, formatos numéricos
AZUL, GRIS, VERDE, NARANJO = '1F3A5F', '5B6670', '2E6B4F', 'A3561A'
bloques = [('MANZENT', 'n_vp', AZUL), (INSUMOS[0], INSUMOS[-1], GRIS), (ORDEN_COMP[0], ORDEN_COMP[-1], VERDE),
           ('z_' + ORDEN_COMP[0], 'z_' + ORDEN_COMP[-1], VERDE), ('n_comp', 'q_urb', NARANJO)]
for a, b, color in bloques:
    for col in range(ws[L[a] + '1'].column, ws[L[b] + '1'].column + 1):
        cell = ws.cell(1, col)
        cell.fill = PatternFill('solid', fgColor=color)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = 13
ws.column_dimensions['A'].width = 17
ws.column_dimensions['B'].width = 13
ws.column_dimensions['C'].width = 15
ws.column_dimensions['D'].width = 16
ws.column_dimensions['E'].width = 16
ws.row_dimensions[1].height = 75
for k in ORDEN_COMP:
    for r in range(R0, R1 + 1):
        ws[ref(k, r)].number_format = '0.0'
for k in ['z_' + k for k in ORDEN_COMP]:
    for r in range(R0, R1 + 1):
        ws[ref(k, r)].number_format = '0.000'
for r in range(R0, R1 + 1):
    ws[ref('indice', r)].number_format = '0.00'
    ws[ref('percentil', r)].number_format = '0.0'
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f"A1:{L['q_urb']}{R1}"

# ---- Parámetros (fórmulas sobre la hoja Manzanas)
wp = wb.create_sheet('Parámetros')
wp.append(['Componente', 'Clave', 'Signo en el índice', 'Media (μ)', 'Desviación estándar (σ)', 'Polígonos con dato', 'Columna en Manzanas'])
for i, (k, v) in enumerate(COMP.items()):
    r = 2 + i
    wp.append([v[2], k, v[1], f'=AVERAGE({rng(k)})', f'=_xlfn.STDEV.S({rng(k)})', f'=COUNT({rng(k)})', L[k]])
    wp[f'D{r}'].number_format = '0.0000'
    wp[f'E{r}'].number_format = '0.0000'
wp['A11'] = 'Mínimo de componentes con dato para calcular el índice'
wp['C11'] = MIN_COMP
wp['A11'].font = Font(bold=True)
wp['A13'] = ('Media y desviación estándar muestral de cada componente sobre todos los polígonos de las 4 comunas con dato. '
             'El signo es -1 en los componentes de carencia. Si se cambia un signo o el mínimo de componentes, todo el libro se recalcula.')
for col, w in zip('ABCDEFG', [44, 10, 12, 14, 16, 14, 14]):
    wp.column_dimensions[col].width = w

# ---- Cortes de quintiles (fórmulas, mismo algoritmo que el mapa)
wc = wb.create_sheet('Cortes de quintiles')
wc.append(['Comuna', 'Alcance', 'Polígonos con índice (n)', 'Mínimo', 'Corte 1', 'Corte 2', 'Corte 3', 'Corte 4',
           'Corte 1 válido', 'Corte 2 válido', 'Corte 3 válido', 'Corte 4 válido'])
fila = 2
for com in COMUNAS.values():
    for alc in ['Todos los polígonos', 'Solo manzanas urbanas']:
        cond = f'({rng("comuna")}=$A{fila})' + ('' if alc.startswith('Todos') else f'*({rng("tipo")}="Manzana urbana")')
        wc.append([com, alc,
                   f'=SUMPRODUCT({cond}*ISNUMBER({rng("indice")}))',
                   f'=_xlfn.AGGREGATE(15,6,{rng("indice")}/({cond}),1)'] +
                  [f'=_xlfn.AGGREGATE(15,6,{rng("indice")}/({cond}),INT({q}*$C{fila}/5)+1)' for q in range(1, 5)] +
                  [f'=E{fila}>D{fila}'] + [f'=AND({get_column_letter(4 + q)}{fila}>D{fila},{get_column_letter(4 + q)}{fila}<>{get_column_letter(3 + q)}{fila})' for q in range(2, 5)])
        fila += 1
wc['A11'] = ('Corte q = valor del índice en la posición INT(q·n/5)+1 de la comuna, ordenada de menor a mayor. Un corte no es válido si es '
             'igual al mínimo o al corte anterior (no ocurre con este índice). Quintil = 1 + número de cortes válidos menores o iguales al índice.')
for col, w in zip('ABCDEFGHIJKL', [14, 22, 14, 10, 10, 10, 10, 10, 11, 11, 11, 11]):
    wc.column_dimensions[col].width = w

# ---- Resumen por comuna (fórmulas)
wr = wb.create_sheet('Resumen por comuna')
wr.append(['Comuna', 'Tipo de polígono', 'Polígonos con índice', 'Personas', 'Media', 'Mediana', 'Mínimo', 'Máximo'])
fila = 2
for com in COMUNAS.values():
    for tipo in ['Manzana urbana', 'Entidad rural']:
        cond = f'({rng("comuna")}=$A{fila})*({rng("tipo")}=$B{fila})'
        wr.append([com, tipo,
                   f'=SUMPRODUCT({cond}*ISNUMBER({rng("indice")}))',
                   f'=SUMPRODUCT({cond}*ISNUMBER({rng("indice")}),{rng("n_per")})',
                   f'=ROUND(SUMIFS({rng("indice")},{rng("comuna")},$A{fila},{rng("tipo")},$B{fila})/C{fila},2)',
                   f'=ROUND(_xlfn.AGGREGATE(16,6,{rng("indice")}/({cond}),0.5),2)',
                   f'=_xlfn.AGGREGATE(15,6,{rng("indice")}/({cond}),1)',
                   f'=_xlfn.AGGREGATE(14,6,{rng("indice")}/({cond}),1)'])
        fila += 1
wr['A11'] = 'Estadísticas entre los polígonos con índice, sin ponderar por población.'
for col, w in zip('ABCDEFGH', [14, 16, 12, 12, 10, 10, 10, 10]):
    wr.column_dimensions[col].width = w

# ---- Diccionario
DICC = [
    ('Columnas grises (insumos)', 'Datos de la base manzana-entidad del Censo 2024 (INE), tal como vienen. Celda vacía = dato suprimido por el INE (menos de 4 casos). Son los únicos valores escritos a mano; todo lo demás es fórmula.'),
    ('MANZENT', 'Identificador único INE de la manzana urbana o entidad rural (llave para cruzar con la cartografía censal).'),
    ('Tipo de polígono', 'Manzana urbana, o entidad rural (polígono de Thiessen trazado por el INE).'),
    ('n_tipo_viv_* (7 columnas)', 'Su suma es el total de viviendas ocupadas con moradores presentes, denominador de hacinamiento y allegamiento.'),
    ('n_mat_paredes_* (6 columnas)', 'Su suma es el denominador de paredes precarias.'),
    ('Componentes (7 columnas verdes)', 'Fórmulas del documento metodológico, redondeadas a un decimal. Quedan vacías si el numerador está suprimido o el denominador es 0.'),
    ('z (7 columnas verdes)', 'signo × (valor − μ) / σ, con μ, σ y el signo de la hoja Parámetros. Positivo = mejor que el promedio.'),
    ('Componentes con dato', 'Cuántos de los 7 puntajes z tienen valor.'),
    ('Índice socioeconómico', 'Promedio simple de los z disponibles, si hay al menos el mínimo de la hoja Parámetros (4). 0 = promedio de los polígonos de las 4 comunas.'),
    ('Percentil (4 comunas)', '% de polígonos de las 4 comunas con un índice estrictamente menor.'),
    ('Quintil en la comuna (todos los polígonos)', 'Grupo 1 (más bajo) a 5 (más alto) dentro de la comuna, con manzanas y entidades rurales. Es el que muestra el mapa con las entidades rurales encendidas.'),
    ('Quintil en la comuna (solo manzanas urbanas)', 'Igual, calculado solo con manzanas urbanas. Es el que muestra el mapa por defecto. Vacío en las entidades rurales.'),
]
wd = wb.create_sheet('Diccionario')
wd.append(['Columna', 'Descripción'])
for row in DICC:
    wd.append(list(row))
wd.column_dimensions['A'].width = 44
wd.column_dimensions['B'].width = 110

for w in wb.worksheets[1:]:
    for cell in w[1]:
        cell.fill = PatternFill('solid', fgColor=AZUL)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    w.row_dimensions[1].height = 34
    for row in w.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=cell.column == 1 or w.title == 'Diccionario', vertical='top')
for w, celda in [(wp, 'A13'), (wc, 'A11'), (wr, 'A11')]:
    w[celda].alignment = Alignment(wrap_text=False)
    w[celda].font = Font(italic=True, color='5B6670')

os.makedirs(SALIDA_DIR, exist_ok=True)
wb.save(SALIDA)

# ---- que Excel calcule y guarde los valores, para que el libro abra ya calculado
ps = (f"$x=New-Object -ComObject Excel.Application; $x.Visible=$false; $x.DisplayAlerts=$false; "
      f"$b=$x.Workbooks.Open('{os.path.abspath(SALIDA)}'); $x.CalculateFull(); $b.Save(); $b.Close($true); $x.Quit(); 'calculado'")
print(subprocess.run(['powershell.exe', '-NoProfile', '-Command', ps], capture_output=True, text=True).stdout.strip())

# ---- validación: valores calculados por Excel contra el cálculo en Python (y por ende contra el sitio)
wv = load_workbook(SALIDA, data_only=True)
m = wv['Manzanas']
py = {str(int(a)): row for a, row in zip(d['MANZENT'].to_list(), d.select(['indice', 'percentil', 'quintil_comuna', 'quintil_comuna_urbano'] + list(COMP)).rows())}
dif = {'indice': 0, 'percentil': 0, 'q_todos': 0, 'q_urb': 0, 'componentes': 0}
errores = 0
for r in range(R0, R1 + 1):
    fila_xl = {k: m[ref(k, r)].value for k in ['MANZENT', 'indice', 'percentil', 'q_todos', 'q_urb'] + ORDEN_COMP}
    if any(isinstance(v, str) and v.startswith('#') for v in fila_xl.values()):
        errores += 1
    p = py[fila_xl['MANZENT']]
    norm = lambda v: None if v in ('', None) else round(float(v), 4)
    for j, k in enumerate(['indice', 'percentil', 'q_todos', 'q_urb']):
        if norm(fila_xl[k]) != norm(p[j]):
            dif[k] += 1
    if any(norm(fila_xl[k]) != norm(p[4 + j]) for j, k in enumerate(ORDEN_COMP)):
        dif['componentes'] += 1
print('celdas con error de Excel:', errores, '| diferencias Excel vs Python:', dif)
print('Parámetros (Excel):', [(wv['Parámetros'][f'B{r}'].value, round(wv['Parámetros'][f'D{r}'].value, 4), round(wv['Parámetros'][f'E{r}'].value, 4)) for r in range(2, 9)])
print('Excel:', SALIDA)
# El documento metodológico (indice_socioeconomico_manzanas_documento.md, en esta carpeta) copia a mano los
# parámetros, cortes y resumen de las hojas del Excel. Si cambian los datos, actualizarlo y regenerar el Word con:
# quarto pandoc indice_socioeconomico_manzanas_documento.md -o "..\Outputs\Indice socioeconomico manzanas\indice_socioeconomico_manzanas.docx"
