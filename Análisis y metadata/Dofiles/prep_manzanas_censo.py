# Prepara los polígonos de manzanas y entidades rurales del Censo 2024 (INE, geoparquet R10)
# para la pestaña Censo 2024 del mapa del sitio. Salida: un .js con los datos listos para Leaflet,
# que el sitio carga recién cuando se abre esa pestaña. Los colegios los toma el sitio de colegiosGeo.
# Correr con: C:\Users\gonza\AppData\Local\Programs\Python\Python312\python.exe prep_manzanas_censo.py
#
# Cada polígono lleva v = [personas, viviendas particulares, indicador_1, ..., indicador_n], en el orden
# de la lista MZK que también se escribe al .js. El sitio busca cada indicador por su clave en MZK,
# así que agregar o reordenar indicadores aquí no rompe el HTML (solo hay que describirlos allá).
import glob, json, os
import numpy as np
import polars as pl
import shapely
from shapely import wkb

BASE = os.path.join(os.path.dirname(__file__), '..', '..')
CARTO = os.path.join(BASE, 'Datos', 'Censo 2024', 'cartografia')
SALIDA = r'C:\Users\gonza\repos\ceel-estudio-puerto-varas\datos_manzanas.js'

# todas las comunas de la Región de Los Lagos (el geoparquet es de la región 10); el nombre se arma desde COMUNA del INE
COMUNAS = {}
# Tolerancia de simplificación en grados (~2 m en manzanas, ~40 m en entidades rurales; con las 30 comunas
# las entidades rurales eran la mitad del peso del archivo)
TOL = {'Manzanas': 0.00002, 'Entidades': 0.0004}

c = pl.col


def s(*cols):
    # suma de categorías; las suprimidas (null, <4 casos) cuentan como 0 en el denominador
    return pl.sum_horizontal([c(x).fill_null(0) for x in cols])


def tasa(num, den):
    # null si el numerador fue suprimido por el INE (<4 casos) o no hay denominador. El redondeo previo a 10 decimales
    # quita el ruido de punto flotante, igual que ROUND de Excel (12,35 guardado como 12,3499999 queda en 12,4)
    return pl.when(num.is_not_null() & den.is_not_null() & (den > 0)).then((num / den * 100).round(10).round(1, mode='half_away_from_zero')).otherwise(None)


def suma_o_null(*cols):
    # numerador compuesto: null solo si todas sus partes vienen suprimidas
    return pl.when(pl.any_horizontal([c(x).is_not_null() for x in cols])).then(s(*cols)).otherwise(None)


TENENCIA = ['n_tenencia_propia_pagada', 'n_tenencia_propia_pagandose', 'n_tenencia_arrendada_contrato',
            'n_tenencia_arrendada_sin_contrato', 'n_tenencia_cedida_trabajo', 'n_tenencia_cedida_familiar', 'n_tenencia_otro']
TIPO_VIV = ['n_tipo_viv_casa', 'n_tipo_viv_depto', 'n_tipo_viv_indigena', 'n_tipo_viv_pieza', 'n_tipo_viv_mediagua',
            'n_tipo_viv_movil', 'n_tipo_viv_otro']
DORM = ['n_dormitorios_1', 'n_dormitorios_2', 'n_dormitorios_3', 'n_dormitorios_4', 'n_dormitorios_5', 'n_dormitorios_6_o_mas']
PAREDES = ['n_mat_paredes_hormigon', 'n_mat_paredes_albanileria', 'n_mat_paredes_tabique_forrado',
           'n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios']
SERV_HIG = ['n_serv_hig_alc_dentro', 'n_serv_hig_alc_fuera', 'n_serv_hig_fosa', 'n_serv_hig_pozo', 'n_serv_hig_acequia_canal',
            'n_serv_hig_cajon_otro', 'n_serv_hig_bano_quimico', 'n_serv_hig_bano_seco', 'n_serv_hig_no_tiene']
CALEF = ['n_comb_calefaccion_gas', 'n_comb_calefaccion_parafina', 'n_comb_calefaccion_lena', 'n_comb_calefaccion_pellet',
         'n_comb_calefaccion_carbon', 'n_comb_calefaccion_electricidad', 'n_comb_calefaccion_otra', 'n_comb_calefaccion_no_utiliza']
ESTCIV = ['n_estcivcon_casado', 'n_estcivcon_conviviente', 'n_estcivcon_conv_civil', 'n_estcivcon_anul_sep_div',
          'n_estcivcon_viudo', 'n_estcivcon_soltero']

menores18 = s('n_edad_0_5', 'n_edad_6_13', 'n_edad_14_17')
mayores18 = c('n_per') - menores18
viv_ocup = s(*TIPO_VIV)  # viviendas particulares ocupadas con moradores presentes
ocupados = c('n_ocupado')

# clave -> expresión. El orden define la posición en v (después de personas y viviendas).
IND = {
    # Población
    'edad': c('prom_edad').round(1, mode='half_away_from_zero'),
    'men18': tasa(suma_o_null('n_edad_0_5', 'n_edad_6_13', 'n_edad_14_17'), c('n_per')),
    'men6': tasa(c('n_edad_0_5'), c('n_per')),
    'may60': tasa(c('n_edad_60_mas'), c('n_per')),
    'inm': tasa(c('n_inmigrantes'), c('n_per')),
    # NNA por tramo de edad del INE (conteos): con el área del polígono, el sitio arma NNA por hectárea para los tramos elegidos
    'n05': c('n_edad_0_5'),
    'n613': c('n_edad_6_13'),
    'n1417': c('n_edad_14_17'),
    # Hogares y viviendas
    'perhog': c('prom_per_hog').round(2, mode='half_away_from_zero'),
    'prop': tasa(suma_o_null('n_tenencia_propia_pagada', 'n_tenencia_propia_pagandose'), s(*TENENCIA)),
    'arr': tasa(suma_o_null('n_tenencia_arrendada_contrato', 'n_tenencia_arrendada_sin_contrato'), s(*TENENCIA)),
    'desoc': tasa(c('n_vp_desocupada'), c('n_vp')),
    'depto': tasa(c('n_tipo_viv_depto'), viv_ocup),
    'dorm4': tasa(suma_o_null('n_dormitorios_4', 'n_dormitorios_5', 'n_dormitorios_6_o_mas'), s(*DORM)),
    'hacin': tasa(c('n_viv_hacinadas'), viv_ocup),
    'alleg': tasa(c('n_hog_allegados'), viv_ocup),
    'irrec': tasa(c('n_viv_irrecuperables'), viv_ocup),
    'precar': tasa(suma_o_null('n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios'), s(*PAREDES)),
    'alcant': tasa(c('n_serv_hig_alc_dentro'), s(*SERV_HIG)),
    'lena': tasa(c('n_comb_calefaccion_lena'), s(*CALEF)),
    # Socioeconómico
    'esc': c('prom_escolaridad18').round(1, mode='half_away_from_zero'),
    'sup': tasa(c('n_cine_terciaria_maestria_doctorado'), mayores18),
    'prof': tasa(suma_o_null('n_ciuo_1', 'n_ciuo_2'), ocupados),
    'elem': tasa(c('n_ciuo_9'), ocupados),
    'desocup': tasa(c('n_desocupado'), s('n_ocupado', 'n_desocupado')),
    'compu': tasa(c('n_serv_compu'), c('n_hog')),
    'inet': tasa(c('n_serv_internet_fija'), c('n_hog')),
    # Composición del hogar
    'casado': tasa(c('n_estcivcon_casado'), s(*ESTCIV)),
    'conviv': tasa(c('n_estcivcon_conviviente'), s(*ESTCIV)),
    'auc': tasa(c('n_estcivcon_conv_civil'), s(*ESTCIV)),
    'sepdiv': tasa(c('n_estcivcon_anul_sep_div'), s(*ESTCIV)),
    'viudo': tasa(c('n_estcivcon_viudo'), s(*ESTCIV)),
    'soltero': tasa(c('n_estcivcon_soltero'), s(*ESTCIV)),
    'unip': tasa(c('n_hog_unipersonales'), c('n_hog')),
    'hog60': tasa(c('n_hog_60'), c('n_hog')),
    'hogmen': tasa(c('n_hog_menores'), c('n_hog')),
    'jefm': tasa(c('n_jefatura_mujer'), c('n_hog')),
}

# Índice socioeconómico (proxy, en la línea del ISMT del Observatorio de Ciudades UC): promedio de
# puntajes z de 7 componentes, con signo negativo en los que indican carencia. Los z usan la media y la
# desviación estándar de la comuna del polígono, ponderadas por su población de 18+; se exige al menos 4
# componentes con dato y 4 personas de 18+. Misma definición que indice_socioeconomico_manzanas_ponderado.py
# Un componente se estandariza en una comuna solo si su n efectivo de Kish, (Σw)²/Σw², llega a MIN_NEFF; si no, no entra
# al índice de esa comuna (con 30 quedan sin índice Curaco de Vélez y Palena).
ISE = {'esc': 1, 'sup': 1, 'prof': 1, 'compu': 1, 'hacin': -1, 'precar': -1, 'alleg': -1}
MIN_COMP = 4
MIN_POB18 = 4
MIN_NEFF = 30
EXCEL_PONDERADO = os.path.join(BASE, 'Análisis y metadata', 'Outputs', 'Indice socioeconomico manzanas',
                               'indice_socioeconomico_manzanas_censo2024_ponderado.xlsx')


def rnd(xy):
    if isinstance(xy[0], (list, tuple)):
        return [rnd(x) for x in xy]
    return [round(xy[0], 5), round(xy[1], 5)]


partes = []
for capa in ['Manzanas', 'Entidades']:
    f = glob.glob(os.path.join(CARTO, f'*_{capa}.parquet'))[0]
    d = pl.read_parquet(f).with_columns(c('CUT').cast(pl.Int64))
    d = d.filter(c('MZ_BASE_CENSO') == 1)
    for cut, nom in d.select('CUT', 'COMUNA').unique().iter_rows():
        t = nom.title()
        for x in [' De ', ' La ', ' Del ']:
            t = t.replace(x, x.lower())
        COMUNAS[cut] = t
    num = [x for x in d.columns if x.startswith(('n_', 'prom_'))]
    d = d.with_columns([c(x).cast(pl.Float64) for x in num])
    d = d.with_columns(**IND).with_columns(capa=pl.lit(capa), pob18=mayores18)
    partes.append(d.select(['capa', 'CUT', 'MANZENT', 'ENTIDAD', 'LOCALIDAD', 'SHAPE', 'n_per', 'n_vp', 'pob18'] + list(IND)))
d = pl.concat(partes)

# área de cada polígono en hectáreas, sobre la geometría original (antes de simplificar). Las coordenadas vienen en
# grados (SIRGAS 2000): se pasan a metros con la proyección sinusoidal centrada en 73° O, que conserva las áreas
# (x = R·Δlon·cos(lat), y = R·lat, con R = radio auténtico de la Tierra); 1 ha = 10.000 m².
R = 6371007.2


def sinusoidal(xy):
    lon, lat = np.radians(xy[:, 0] + 73), np.radians(xy[:, 1])
    return np.column_stack([R * lon * np.cos(lat), R * lat])


area_ha = shapely.area(shapely.transform(shapely.from_wkb(d['SHAPE'].to_numpy()), sinusoidal)) / 10000
# 4 decimales (1 m²): las manzanas chicas tienen menos de media hectárea. El sitio calcula NNA por hectárea
# sumando los tramos de edad que se elijan (n05, n613, n1417) y dividiendo por ha
d = d.with_columns(ha=pl.Series(area_ha).round(4))

# ponderador: población de 18+ (0 si no llega al mínimo); el sitio lo usa para los quintiles ponderados
d = d.with_columns(w=pl.when(c('pob18') >= MIN_POB18).then(c('pob18')).otherwise(0.0))
z = []
for k, sgn in ISE.items():
    con = c(k).is_not_null() & (c('w') > 0)
    sw = c('w').filter(con).sum().over('CUT')
    mu = (c('w') * c(k)).filter(con).sum().over('CUT') / sw
    de = (((c('w') * c(k) * c(k)).filter(con).sum().over('CUT') - sw * mu * mu) / (sw - 1)).sqrt()
    neff = sw * sw / (c('w') * c('w')).filter(con).sum().over('CUT')
    z.append(pl.when(con & (neff >= MIN_NEFF) & (de > 0)).then(sgn * (c(k) - mu) / de).otherwise(None).alias('z_' + k))
d = d.with_columns(z)
zc = ['z_' + k for k in ISE]
d = d.with_columns(
    ise=pl.when((c('w') > 0) & (pl.sum_horizontal([c(x).is_not_null() for x in zc]) >= MIN_COMP))
    .then(pl.mean_horizontal(zc).round(2, mode='half_away_from_zero')).otherwise(None)
)
d = d.with_columns(c('w').cast(pl.Int64), c('n05').cast(pl.Int64), c('n613').cast(pl.Int64), c('n1417').cast(pl.Int64))
MZK = list(IND) + ['ha', 'ise', 'w']

# control: el índice tiene que calzar con el Excel ponderado (columna "Índice 7 componentes")
if os.path.exists(EXCEL_PONDERADO):
    from openpyxl import load_workbook
    hoja = load_workbook(EXCEL_PONDERADO, read_only=True, data_only=True)['Manzanas']
    enc = next(hoja.iter_rows(max_row=1, values_only=True))
    iid, iind = enc.index('MANZENT'), enc.index('Índice 7 componentes')
    xl = {int(r[iid]): (None if r[iind] in ('', None) else round(r[iind], 2)) for r in hoja.iter_rows(min_row=2, values_only=True)}
    dif = sum(1 for a, b in zip(d['MANZENT'].to_list(), d['ise'].to_list()) if xl.get(int(a)) != b)
    print('índice: diferencias contra el Excel ponderado:', dif, 'de', d.height)

filas = []
for r in d.iter_rows(named=True):
    capa = r['capa']
    g = shapely.simplify(wkb.loads(r['SHAPE']), TOL[capa], preserve_topology=True)
    gj = shapely.geometry.mapping(g)
    nombre = r['ENTIDAD'] if capa == 'Entidades' else r['LOCALIDAD']
    filas.append({
        't': 'M' if capa == 'Manzanas' else 'E',
        'c': COMUNAS[r['CUT']],
        'id': str(int(r['MANZENT'])),
        'nom': (nombre or '').title(),
        'g': {'type': gj['type'], 'coordinates': rnd(gj['coordinates'])},
        'v': [int(r['n_per'] or 0), int(r['n_vp'] or 0)] + [r[k] for k in MZK],
    })

with open(SALIDA, 'w', encoding='utf8') as f:
    f.write('var MZK=' + json.dumps(MZK) + ';\n')
    f.write('var MZ=' + json.dumps(filas, ensure_ascii=False, separators=(',', ':')) + ';\n')
print(len(filas), 'poligonos,', len(MZK), 'indicadores,', round(os.path.getsize(SALIDA) / 1e6, 2), 'MB')

# control rápido: cobertura de cada indicador (polígonos con dato) y rango
res = d.select([pl.struct(n=c(k).is_not_null().sum(), mn=c(k).min(), md=c(k).median(), mx=c(k).max()).alias(k) for k in MZK]).row(0, named=True)
for k in MZK:
    print(f"  {k:8s} con dato {res[k]['n']:5d}  min {res[k]['mn']}  mediana {res[k]['md']}  max {res[k]['mx']}")
