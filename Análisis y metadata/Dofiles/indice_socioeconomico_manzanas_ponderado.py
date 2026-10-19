# Índice socioeconómico (proxy) por manzana, Censo 2024, versión ponderada e intracomunal: Excel con fórmulas vivas.
# Diferencias con indice_socioeconomico_manzanas.py (la versión que hoy está en el sitio):
#   - cubre las 30 comunas de la Región de Los Lagos;
#   - la media y la desviación estándar de cada componente se calculan dentro de cada comuna;
#   - todo se pondera por la población de 18 años o más del polígono (media, desviación, percentil, quintiles, correlaciones);
#   - tres indicadores para probar robustez: índice de 7 componentes, índice de 6 (sin escolaridad) y escolaridad promedio 18+.
# No toca el sitio (index.html ni datos_manzanas.js).
# Correr con: C:\Users\gonza\AppData\Local\Programs\Python\Python312\python.exe indice_socioeconomico_manzanas_ponderado.py
import glob, math, os, subprocess
import numpy as np
import polars as pl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
CARTO = os.path.join(BASE, 'Datos', 'Censo 2024', 'cartografia')
PERSONAS = os.path.join(BASE, 'Datos', 'Censo 2024', 'personas_censo2024.parquet')
SALIDA_DIR = os.path.join(BASE, 'Análisis y metadata', 'Outputs', 'Indice socioeconomico manzanas')
SALIDA = os.environ.get('SALIDA_XLSX') or os.path.join(SALIDA_DIR, 'indice_socioeconomico_manzanas_censo2024_ponderado.xlsx')
c = pl.col

# criterios (quedan como celdas editables en la hoja Criterios)
MIN_COMP = 4    # componentes con dato para calcular el índice
MIN_POB18 = 4   # población de 18+ mínima para ponderar el polígono
MIN_NEFF = 30   # tamaño efectivo de Kish mínimo para estandarizar un componente dentro de una comuna
TOL_ESC = 0.05  # diferencia máxima aceptada (años) entre la escolaridad ponderada y la del microdato


def s(*cols):
    return pl.sum_horizontal([c(x).fill_null(0) for x in cols])


def tasa(num, den):
    # el redondeo previo a 10 decimales quita el ruido de punto flotante, igual que ROUND de Excel (sin él, un 12,35 guardado
    # como 12,3499999 se redondea a 12,3 en Python y a 12,4 en Excel)
    return pl.when(num.is_not_null() & den.is_not_null() & (den > 0)).then((num / den * 100).round(10).round(1, mode='half_away_from_zero')).otherwise(None)


def suma_o_null(*cols):
    return pl.when(pl.any_horizontal([c(x).is_not_null() for x in cols])).then(s(*cols)).otherwise(None)


TIPO_VIV = ['n_tipo_viv_casa', 'n_tipo_viv_depto', 'n_tipo_viv_indigena', 'n_tipo_viv_pieza', 'n_tipo_viv_mediagua',
            'n_tipo_viv_movil', 'n_tipo_viv_otro']
PAREDES = ['n_mat_paredes_hormigon', 'n_mat_paredes_albanileria', 'n_mat_paredes_tabique_forrado',
           'n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios']
mayores18 = c('n_per') - s('n_edad_0_5', 'n_edad_6_13', 'n_edad_14_17')
viv_ocup = s(*TIPO_VIV)

# componente: (expresión, signo, nombre para el Excel); mismas definiciones que la versión del sitio
COMP = {
    'esc': (c('prom_escolaridad18').round(1, mode='half_away_from_zero'), 1, 'Escolaridad promedio 18+ (años)'),
    'sup': (tasa(c('n_cine_terciaria_maestria_doctorado'), mayores18), 1, '% educación superior (18+)'),
    'prof': (tasa(suma_o_null('n_ciuo_1', 'n_ciuo_2'), c('n_ocupado')), 1, '% directivos y profesionales (ocupados)'),
    'compu': (tasa(c('n_serv_compu'), c('n_hog')), 1, '% hogares con computador'),
    'hacin': (tasa(c('n_viv_hacinadas'), viv_ocup), -1, '% viviendas hacinadas'),
    'precar': (tasa(suma_o_null('n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios'), s(*PAREDES)), -1, '% viviendas con paredes precarias'),
    'alleg': (tasa(c('n_hog_allegados'), viv_ocup), -1, 'Hogares allegados por 100 viviendas'),
}
ORDEN_COMP = list(COMP)
COMP6 = ORDEN_COMP[1:]  # índice sin el componente de escolaridad

INSUMOS = ['prom_escolaridad18', 'n_per', 'n_edad_0_5', 'n_edad_6_13', 'n_edad_14_17', 'n_cine_terciaria_maestria_doctorado',
           'n_ocupado', 'n_ciuo_1', 'n_ciuo_2', 'n_hog', 'n_serv_compu', 'n_viv_hacinadas', 'n_hog_allegados'] + TIPO_VIV + PAREDES


def nombre_comuna(x):
    t = x.title()
    for p in [' De ', ' La ', ' Del ']:
        t = t.replace(p, p.lower())
    return t


partes = []
for capa in ['Manzanas', 'Entidades']:
    f = glob.glob(os.path.join(CARTO, f'*_{capa}.parquet'))[0]
    d = pl.read_parquet(f).drop('SHAPE').with_columns(c('CUT').cast(pl.Int64))
    d = d.filter(c('MZ_BASE_CENSO') == 1)
    num = [x for x in d.columns if x.startswith(('n_', 'prom_'))]
    d = d.with_columns([c(x).cast(pl.Float64) for x in num])
    d = d.with_columns(**{k: v[0] for k, v in COMP.items()})
    d = d.with_columns(tipo=pl.lit('Manzana urbana' if capa == 'Manzanas' else 'Entidad rural'))
    partes.append(d.select(['MANZENT', 'CUT', 'COMUNA', 'tipo', 'LOCALIDAD', 'ENTIDAD', 'n_vp'] + INSUMOS + ORDEN_COMP))
d = pl.concat(partes).with_columns(comuna=c('COMUNA').map_elements(nombre_comuna, return_dtype=pl.Utf8)).drop('COMUNA')
d = d.sort(['CUT', 'tipo', 'MANZENT'], descending=[False, True, False])  # Manzana urbana antes que Entidad rural
CUTS = sorted(d['CUT'].unique().to_list())
NOMBRE = dict(d.select('CUT', 'comuna').unique().iter_rows())

# ponderador: población de 18 años o más; menos de MIN_POB18 personas = no pondera (peso 0)
d = d.with_columns(pob18=c('n_per') - (c('n_edad_0_5') + c('n_edad_6_13') + c('n_edad_14_17')))
d = d.with_columns(w=pl.when(c('pob18') >= MIN_POB18).then(c('pob18')).otherwise(0.0))

# ---- parámetros ponderados por comuna y componente (mismas sumas que hace el Excel)
par = {}
for cut in CUTS:
    sub = d.filter(c('CUT') == cut)
    for k in ORDEN_COMP:
        x = sub[k].to_numpy()
        w = sub['w'].to_numpy()
        ok = ~np.isnan(x)
        sw, swx, swxx, sww = w[ok].sum(), (w[ok] * x[ok]).sum(), (w[ok] * x[ok] * x[ok]).sum(), (w[ok] * w[ok]).sum()
        mu = swx / sw if sw > 0 else None
        de = math.sqrt((swxx - sw * mu * mu) / (sw - 1)) if sw > 1 else None
        neff = sw * sw / sww if sww > 0 else 0
        par[(cut, k)] = {'sw': sw, 'mu': mu, 'de': de, 'neff': neff, 'valido': bool(neff >= MIN_NEFF and de and de > 0)}

zs = []
for k in ORDEN_COMP:
    sgn = COMP[k][1]
    mu = c('CUT').replace_strict({cut: par[(cut, k)]['mu'] for cut in CUTS}, return_dtype=pl.Float64)
    de = c('CUT').replace_strict({cut: par[(cut, k)]['de'] for cut in CUTS}, return_dtype=pl.Float64)
    val = c('CUT').replace_strict({cut: par[(cut, k)]['valido'] for cut in CUTS}, return_dtype=pl.Boolean)
    zs.append(pl.when(c(k).is_not_null() & (c('w') > 0) & val).then(sgn * (c(k) - mu) / de).otherwise(None).alias('z_' + k))
d = d.with_columns(zs)
z7 = ['z_' + k for k in ORDEN_COMP]
z6 = ['z_' + k for k in COMP6]
d = d.with_columns(n_comp7=pl.sum_horizontal([c(x).is_not_null() for x in z7]).cast(pl.Int8),
                   n_comp6=pl.sum_horizontal([c(x).is_not_null() for x in z6]).cast(pl.Int8))
d = d.with_columns(
    i7=pl.when((c('w') > 0) & (c('n_comp7') >= MIN_COMP)).then(pl.mean_horizontal(z7).round(2, mode='half_away_from_zero')).otherwise(None),
    i6=pl.when((c('w') > 0) & (c('n_comp6') >= MIN_COMP)).then(pl.mean_horizontal(z6).round(2, mode='half_away_from_zero')).otherwise(None),
    iesc=pl.when(c('w') > 0).then(c('esc')).otherwise(None),
)
d = d.with_columns([c(x).round(3, mode='half_away_from_zero') for x in z7])  # solo para mostrar
d = d.with_columns(estado=pl.when(c('pob18') < MIN_POB18).then(pl.lit('Sin población de 18+ suficiente para ponderar'))
                   .when(c('n_comp7') < MIN_COMP).then(pl.lit('Menos de 4 componentes con dato'))
                   .otherwise(pl.lit('Con índice')))

# ---- percentil y quintiles ponderados dentro de la comuna: % de la población 18+ en polígonos con índice estrictamente menor
IDX = {'i7': 'Índice 7 componentes', 'i6': 'Índice 6 componentes (sin escolaridad)', 'iesc': 'Escolaridad promedio 18+'}


def pct_ponderado(vals, w):
    out = np.full(len(vals), np.nan)
    ok = ~np.isnan(vals)
    if not ok.any():
        return out
    v, ww = vals[ok], w[ok]
    orden = np.argsort(v, kind='stable')
    vs, cw = v[orden], np.concatenate([[0.0], np.cumsum(ww[orden])])
    den = cw[-1]
    num = cw[np.searchsorted(vs, v, side='left')]
    out[ok] = 100 * num / den
    return out


extra = {f'{k}_{t}': np.full(d.height, np.nan) for k in IDX for t in ['pct', 'q', 'qurb']}
cut_arr, tipo_arr, w_arr = d['CUT'].to_numpy(), d['tipo'].to_numpy(), d['w'].to_numpy()
for k in IDX:
    vals = d[k].to_numpy().astype(float)
    for cut in CUTS:
        m = cut_arr == cut
        p = pct_ponderado(vals[m], w_arr[m])
        extra[f'{k}_pct'][m] = p
        extra[f'{k}_q'][m] = np.where(np.isnan(p), np.nan, np.minimum(5, 1 + np.floor(p / 20)))
        mu_ = m & (tipo_arr == 'Manzana urbana')
        vu = vals[mu_]
        ok = ~np.isnan(vu)
        if ok.any():
            orden = np.argsort(vu[ok], kind='stable')
            vs, cw = vu[ok][orden], np.concatenate([[0.0], np.cumsum(w_arr[mu_][ok][orden])])
            q = np.full(len(vu), np.nan)
            q[ok] = np.minimum(5, 1 + np.floor(5 * cw[np.searchsorted(vs, vu[ok], side='left')] / cw[-1]))
            extra[f'{k}_qurb'][mu_] = q
d = d.with_columns([pl.Series(n, v).fill_nan(None) for n, v in extra.items()])

# ---- microdato de personas para la validación (escolaridad -99 = sin dato)
mic = (pl.scan_parquet(PERSONAS).filter((c('region') == 10) & (c('edad') >= 18)).group_by('comuna')
       .agg(pob18_mic=pl.len(), esc_mic=c('escolaridad').filter(c('escolaridad') >= 0).mean()).collect())
MIC = {int(r['comuna']): r for r in mic.iter_rows(named=True)}

# ---- resultados en consola
print('polígonos:', d.height, '| comunas:', len(CUTS), '| sin población 18+ suficiente:', d.filter(c('pob18') < MIN_POB18).height,
      '| con índice 7:', d['i7'].drop_nulls().len(), '| índice 6:', d['i6'].drop_nulls().len(), '| escolaridad:', d['iesc'].drop_nulls().len())
inval = [(NOMBRE[cut], k, round(par[(cut, k)]['neff'], 1)) for cut in CUTS for k in ORDEN_COMP if not par[(cut, k)]['valido']]
print('componentes no estandarizados (n efectivo < %d):' % MIN_NEFF, inval)


def corr_w(a, b, w):
    ok = ~np.isnan(a) & ~np.isnan(b)
    a, b, w = a[ok], b[ok], w[ok]
    sw = w.sum()
    ma, mb = (w * a).sum() / sw, (w * b).sum() / sw
    return ((w * (a - ma) * (b - mb)).sum()) / math.sqrt((w * (a - ma) ** 2).sum() * (w * (b - mb) ** 2).sum())


PARES = [('i7', 'i6'), ('i7', 'iesc'), ('i6', 'iesc')]
for cut in [10109, 10101, 10105, 10301]:
    sub = d.filter(c('CUT') == cut)
    print(NOMBRE[cut], {f'{a}-{b}': round(corr_w(sub[a].to_numpy().astype(float), sub[b].to_numpy().astype(float), sub['w'].to_numpy()), 3) for a, b in PARES})

# ==== Excel con fórmulas vivas
N = d.height
R0, R1 = 2, N + 1
L = {}
NOMBRES_COL = (['MANZENT', 'CUT', 'comuna', 'tipo', 'loc', 'ent', 'n_vp'] + INSUMOS + ['pob18', 'w', 'fila_par'] + ORDEN_COMP + z7
               + ['n_comp7', 'i7', 'n_comp6', 'i6', 'iesc', 'estado']
               + [f'w_{k}' for k in IDX] + [f'{k}_pct' for k in IDX] + [f'{k}_q' for k in IDX] + [f'{k}_qurb' for k in IDX])
for i, nombre in enumerate(NOMBRES_COL, start=1):
    L[nombre] = get_column_letter(i)
ETIQ = {'MANZENT': 'MANZENT', 'CUT': 'CUT', 'comuna': 'Comuna', 'tipo': 'Tipo de polígono', 'loc': 'Localidad', 'ent': 'Entidad',
        'n_vp': 'n_vp (viviendas particulares)', 'pob18': 'Población 18+ (ponderador)', 'w': 'Peso usado (0 si población 18+ < mínimo)',
        'fila_par': 'Fila de la comuna en Parámetros', 'n_comp7': 'Componentes con dato (de 7)', 'i7': 'Índice 7 componentes',
        'n_comp6': 'Componentes con dato (de 6)', 'i6': 'Índice 6 componentes (sin escolaridad)', 'iesc': 'Índice de escolaridad 18+ (años)',
        'estado': 'Estado del polígono'}
ETIQ.update({k: COMP[k][2] for k in ORDEN_COMP})
ETIQ.update({'z_' + k: 'z ' + COMP[k][2] for k in ORDEN_COMP})
for k, nom in IDX.items():
    ETIQ[f'w_{k}'] = f'Peso en percentil: {nom}'
    ETIQ[f'{k}_pct'] = f'Percentil ponderado en la comuna: {nom}'
    ETIQ[f'{k}_q'] = f'Quintil ponderado en la comuna (todos los polígonos): {nom}'
    ETIQ[f'{k}_qurb'] = f'Quintil ponderado en la comuna (solo manzanas urbanas): {nom}'


def rng(k):
    return f"Manzanas!${L[k]}${R0}:${L[k]}${R1}"


def ref(k, r):
    return f"{L[k]}{r}"


P0 = 3  # primera fila de datos en Parámetros (dos filas de encabezado)
P1 = P0 + len(CUTS) - 1
PCOL = {}  # (componente, dato) -> letra de columna en Parámetros
col = 5
for k in ORDEN_COMP:
    for dato in ['sw', 'mu', 'de', 'neff', 'valido']:
        PCOL[(k, dato)] = get_column_letter(col)
        col += 1
PCOL_VALIDOS = get_column_letter(col)
CR = {'min_comp': 'Criterios!$C$2', 'min_pob': 'Criterios!$C$3', 'min_neff': 'Criterios!$C$4', 'tol': 'Criterios!$C$5'}
SIGNO = {k: f'Criterios!$C${9 + i}' for i, k in enumerate(ORDEN_COMP)}


def formulas_fila(r):
    f = {}
    x = lambda k: ref(k, r)
    f['pob18'] = f'={x("n_per")}-({x("n_edad_0_5")}+{x("n_edad_6_13")}+{x("n_edad_14_17")})'
    f['w'] = f'=IF({x("pob18")}>={CR["min_pob"]},{x("pob18")},0)'
    f['fila_par'] = f"=MATCH({x('CUT')},'Parámetros'!$A${P0}:$A${P1},0)"
    f['esc'] = f'=IF({x("prom_escolaridad18")}="","",ROUND({x("prom_escolaridad18")},1))'
    f['sup'] = f'=IF(OR({x("n_cine_terciaria_maestria_doctorado")}="",{x("pob18")}<=0),"",ROUND(100*{x("n_cine_terciaria_maestria_doctorado")}/{x("pob18")},1))'
    f['prof'] = f'=IF(OR(AND({x("n_ciuo_1")}="",{x("n_ciuo_2")}=""),N({x("n_ocupado")})<=0),"",ROUND(100*({x("n_ciuo_1")}+{x("n_ciuo_2")})/{x("n_ocupado")},1))'
    f['compu'] = f'=IF(OR({x("n_serv_compu")}="",N({x("n_hog")})<=0),"",ROUND(100*{x("n_serv_compu")}/{x("n_hog")},1))'
    viv = f'SUM({ref(TIPO_VIV[0], r)}:{ref(TIPO_VIV[-1], r)})'
    f['hacin'] = f'=IF(OR({x("n_viv_hacinadas")}="",{viv}<=0),"",ROUND(100*{x("n_viv_hacinadas")}/{viv},1))'
    par_ = f'SUM({ref(PAREDES[0], r)}:{ref(PAREDES[-1], r)})'
    a, b, cc = (ref(k, r) for k in ['n_mat_paredes_tabique_sin_forro', 'n_mat_paredes_artesanal', 'n_mat_paredes_precarios'])
    f['precar'] = f'=IF(OR(AND({a}="",{b}="",{cc}=""),{par_}<=0),"",ROUND(100*({a}+{b}+{cc})/{par_},1))'
    f['alleg'] = f'=IF(OR({x("n_hog_allegados")}="",{viv}<=0),"",ROUND(100*{x("n_hog_allegados")}/{viv},1))'
    for k in ORDEN_COMP:
        g = lambda dato: f"INDEX('Parámetros'!${PCOL[(k, dato)]}${P0}:${PCOL[(k, dato)]}${P1},{x('fila_par')})"
        f['z_' + k] = f'=IF(OR({x(k)}="",{x("w")}=0),"",IF({g("valido")},{SIGNO[k]}*({x(k)}-{g("mu")})/{g("de")},""))'
    r7 = f'{ref("z_" + ORDEN_COMP[0], r)}:{ref("z_" + ORDEN_COMP[-1], r)}'
    r6 = f'{ref("z_" + COMP6[0], r)}:{ref("z_" + COMP6[-1], r)}'
    f['n_comp7'] = f'=COUNT({r7})'
    f['i7'] = f'=IF(AND({x("w")}>0,{x("n_comp7")}>={CR["min_comp"]}),ROUND(AVERAGE({r7}),2),"")'
    f['n_comp6'] = f'=COUNT({r6})'
    f['i6'] = f'=IF(AND({x("w")}>0,{x("n_comp6")}>={CR["min_comp"]}),ROUND(AVERAGE({r6}),2),"")'
    f['iesc'] = f'=IF(OR({x("w")}=0,{x("esc")}=""),"",{x("esc")})'
    f['estado'] = (f'=IF({x("pob18")}<{CR["min_pob"]},"Sin población de 18+ suficiente para ponderar",'
                   f'IF({x("n_comp7")}<{CR["min_comp"]},"Menos de 4 componentes con dato","Con índice"))')
    for k in IDX:
        f[f'w_{k}'] = f'=IF(ISNUMBER({x(k)}),{x("w")},0)'
        f[f'{k}_pct'] = (f'=IF({x(k)}="","",100*SUMIFS({rng("w_" + k)},{rng("CUT")},{x("CUT")},{rng(k)},"<"&{x(k)})'
                         f'/SUMIFS({rng("w_" + k)},{rng("CUT")},{x("CUT")}))')
        f[f'{k}_q'] = f'=IF({x(k + "_pct")}="","",MIN(5,1+INT({x(k + "_pct")}/20)))'
        f[f'{k}_qurb'] = (f'=IF(OR({x("tipo")}<>"Manzana urbana",{x(k)}=""),"",MIN(5,1+INT(5*SUMIFS({rng("w_" + k)},{rng("CUT")},{x("CUT")},'
                          f'{rng("tipo")},"Manzana urbana",{rng(k)},"<"&{x(k)})/SUMIFS({rng("w_" + k)},{rng("CUT")},{x("CUT")},{rng("tipo")},"Manzana urbana"))))')
    return f


wb = Workbook()
ws = wb.active
ws.title = 'Manzanas'
ws.append([ETIQ.get(k, k) for k in NOMBRES_COL])
for i, fila in enumerate(d.select(['MANZENT', 'CUT', 'comuna', 'tipo', 'LOCALIDAD', 'ENTIDAD', 'n_vp'] + INSUMOS).rows()):
    r = R0 + i
    mz, cut, com, tipo, loc, ent, nvp, *ins = fila
    ws.append([str(int(mz)), int(cut), com, tipo, loc, ent, None if nvp is None else int(nvp)]
              + [None if v is None else (v if k == 'prom_escolaridad18' else int(v)) for k, v in zip(INSUMOS, ins)])
    for k, frm in formulas_fila(r).items():
        ws[ref(k, r)] = frm

AZUL, GRIS, MORADO, VERDE, NARANJO, CAFE = '1F3A5F', '5B6670', '5B3F86', '2E6B4F', 'A3561A', '7A5C2E'
bloques = [('MANZENT', 'n_vp', AZUL), (INSUMOS[0], INSUMOS[-1], GRIS), ('pob18', 'fila_par', MORADO), (ORDEN_COMP[0], z7[-1], VERDE),
           ('n_comp7', 'estado', NARANJO), ('w_i7', 'iesc_qurb', CAFE)]
for a, b, color in bloques:
    for ci in range(ws[L[a] + '1'].column, ws[L[b] + '1'].column + 1):
        cell = ws.cell(1, ci)
        cell.fill = PatternFill('solid', fgColor=color)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(ci)].width = 13
for k, wd in [('MANZENT', 17), ('comuna', 16), ('tipo', 15), ('loc', 16), ('ent', 16), ('estado', 24)]:
    ws.column_dimensions[L[k]].width = wd
ws.row_dimensions[1].height = 90
FMT = {**{k: '0.0' for k in ORDEN_COMP}, **{z: '0.000' for z in z7}, 'i7': '0.00', 'i6': '0.00', 'iesc': '0.0',
       **{f'{k}_pct': '0.0' for k in IDX}}
for k, fmt in FMT.items():
    for r in range(R0, R1 + 1):
        ws[ref(k, r)].number_format = fmt
ws.freeze_panes = 'D2'
ws.auto_filter.ref = f"A1:{L['iesc_qurb']}{R1}"

# ---- Criterios
wcr = wb.create_sheet('Criterios')
wcr.append(['Criterio', '', 'Valor', 'Uso'])
wcr.append(['Mínimo de componentes con dato para calcular el índice', '', MIN_COMP, 'Índices de 7 y de 6 componentes.'])
wcr.append(['Población de 18+ mínima para ponderar el polígono', '', MIN_POB18, 'Con menos, el polígono queda con peso 0 y sin índice (sigue en la base).'])
wcr.append(['Tamaño efectivo mínimo (Kish) para estandarizar un componente en una comuna', '', MIN_NEFF,
            'Si la comuna no lo alcanza en un componente, ese componente no entra al índice de esa comuna.'])
wcr.append(['Tolerancia del chequeo de escolaridad contra el microdato (años)', '', TOL_ESC, 'Hoja Validación.'])
wcr.append([])
wcr.append(['Signo de cada componente en el índice'])
wcr.append(['Componente', 'Clave', 'Signo'])
for k in ORDEN_COMP:
    wcr.append([COMP[k][2], k, COMP[k][1]])
wcr['A17'] = ('Tamaño efectivo de Kish: n_ef = (Σ w)² / Σ w², sobre los polígonos de la comuna con dato en el componente. Es el número de '
              'polígonos que darían la misma precisión si todos pesaran igual. Con n_ef = 30 el error relativo de la desviación estándar '
              'es cercano a 1/√(2·(n_ef−1)) ≈ 13%.')
for colx, wd in zip('ABCD', [70, 10, 10, 90]):
    wcr.column_dimensions[colx].width = wd

# ---- Parámetros por comuna (fórmulas)
wp = wb.create_sheet('Parámetros')
wp['A1'], wp['B1'], wp['C1'], wp['D1'] = 'CUT', 'Comuna', 'Polígonos', 'Población 18+ ponderada (Σ w)'
for k in ORDEN_COMP:
    wp[f'{PCOL[(k, "sw")]}1'] = COMP[k][2]
    wp.merge_cells(f'{PCOL[(k, "sw")]}1:{PCOL[(k, "valido")]}1')
    for dato, et in zip(['sw', 'mu', 'de', 'neff', 'valido'], ['Σ w con dato', 'Media ponderada (μ)', 'Desv. estándar ponderada (σ)', 'n efectivo (Kish)', 'Se estandariza']):
        wp[f'{PCOL[(k, dato)]}2'] = et
wp[f'{PCOL_VALIDOS}1'] = 'Componentes estandarizados en la comuna'
for i, cut in enumerate(CUTS):
    r = P0 + i
    wp[f'A{r}'] = cut
    wp[f'B{r}'] = NOMBRE[cut]
    wp[f'C{r}'] = f'=COUNTIF({rng("CUT")},$A{r})'
    wp[f'D{r}'] = f'=SUMIFS({rng("w")},{rng("CUT")},$A{r})'
    for k in ORDEN_COMP:
        m = f'--({rng("CUT")}=$A{r})'
        sw, mu, de, ne, va = (f'{PCOL[(k, x)]}{r}' for x in ['sw', 'mu', 'de', 'neff', 'valido'])
        wp[sw] = f'=SUMPRODUCT({m},--ISNUMBER({rng(k)}),{rng("w")})'
        wp[mu] = f'=IF({sw}=0,"",SUMPRODUCT({m},{rng("w")},{rng(k)})/{sw})'
        wp[de] = f'=IF({sw}<=1,"",SQRT((SUMPRODUCT({m},{rng("w")},{rng(k)},{rng(k)})-{sw}*{mu}*{mu})/({sw}-1)))'
        wp[ne] = f'=IF({sw}=0,0,{sw}^2/SUMPRODUCT({m},--ISNUMBER({rng(k)}),{rng("w")},{rng("w")}))'
        wp[va] = f'=AND({ne}>={CR["min_neff"]},N({de})>0)'
        for cel, fmt in [(mu, '0.0000'), (de, '0.0000'), (ne, '0.0'), (sw, '#,##0')]:
            wp[cel].number_format = fmt
    wp[f'{PCOL_VALIDOS}{r}'] = '=' + '+'.join(f'--{PCOL[(k, "valido")]}{r}' for k in ORDEN_COMP)
wp[f'A{P1 + 2}'] = ('Media y desviación estándar de cada componente dentro de la comuna, ponderadas por la población de 18+ del polígono '
                    '(columna "Peso usado" de Manzanas), sobre los polígonos con dato. σ usa Σw − 1 en el denominador (pesos de frecuencia).')
wp.column_dimensions['B'].width = 22
wp.column_dimensions['D'].width = 14
for ci in range(5, col + 1):
    wp.column_dimensions[get_column_letter(ci)].width = 12
wp.freeze_panes = 'C3'

# ---- Correlaciones ponderadas entre los tres indicadores
wco = wb.create_sheet('Correlaciones')
wco.append(['CUT', 'Comuna', 'Par', 'Polígonos con ambos', 'Σ w', 'Σ w·a', 'Σ w·b', 'Σ w·a·b', 'Σ w·a²', 'Σ w·b²', 'r ponderado'])
fila = 2
for cut in CUTS:
    for a, b in PARES:
        m = f'--({rng("CUT")}=$A{fila}),--ISNUMBER({rng(a)}),--ISNUMBER({rng(b)})'
        wco.append([cut, NOMBRE[cut], f'{IDX[a]} vs {IDX[b]}',
                    f'=SUMPRODUCT({m})',
                    f'=SUMPRODUCT({m},{rng("w")})',
                    f'=SUMPRODUCT({m},{rng("w")},{rng(a)})',
                    f'=SUMPRODUCT({m},{rng("w")},{rng(b)})',
                    f'=SUMPRODUCT({m},{rng("w")},{rng(a)},{rng(b)})',
                    f'=SUMPRODUCT({m},{rng("w")},{rng(a)},{rng(a)})',
                    f'=SUMPRODUCT({m},{rng("w")},{rng(b)},{rng(b)})',
                    f'=IFERROR((H{fila}-F{fila}*G{fila}/E{fila})/SQRT((I{fila}-F{fila}^2/E{fila})*(J{fila}-G{fila}^2/E{fila})),"")'])
        wco[f'K{fila}'].number_format = '0.000'
        fila += 1
wco[f'A{fila + 1}'] = 'Correlación de Pearson ponderada por la población de 18+, entre los polígonos de la comuna que tienen ambos indicadores.'
for colx, wd in zip('ABCDEFGHIJK', [8, 22, 62, 12, 12, 14, 14, 14, 14, 14, 12]):
    wco.column_dimensions[colx].width = wd
wco.freeze_panes = 'D2'

# ---- Población por quintil (% real de población 18+ en cada quintil)
wq = wb.create_sheet('Población por quintil')
wq.append(['CUT', 'Comuna', 'Indicador', 'Alcance', 'Población 18+ con indicador']
          + [f'% población Q{q}' for q in range(1, 6)] + [f'Polígonos Q{q}' for q in range(1, 6)]
          + [f'Mínimo del indicador en Q{q}' for q in range(1, 6)])
fila = 2
for cut in CUTS:
    for k in IDX:
        for alc, qc in [('Todos los polígonos', f'{k}_q'), ('Solo manzanas urbanas', f'{k}_qurb')]:
            urb = '' if alc.startswith('Todos') else f',{rng("tipo")},"Manzana urbana"'
            fila_v = [cut, NOMBRE[cut], IDX[k], alc, f'=SUMIFS({rng("w_" + k)},{rng("CUT")},$A{fila}{urb})']
            fila_v += [f'=IF($E{fila}=0,"",SUMIFS({rng("w_" + k)},{rng("CUT")},$A{fila},{rng(qc)},{q})/$E{fila})' for q in range(1, 6)]
            fila_v += [f'=COUNTIFS({rng("CUT")},$A{fila},{rng(qc)},{q})' for q in range(1, 6)]
            fila_v += [f'=IF({get_column_letter(10 + q)}{fila}=0,"",_xlfn.MINIFS({rng(k)},{rng("CUT")},$A{fila},{rng(qc)},{q}))' for q in range(1, 6)]
            wq.append(fila_v)
            for q in range(1, 6):
                wq[f'{get_column_letter(5 + q)}{fila}'].number_format = '0.0%'
            wq[f'E{fila}'].number_format = '#,##0'
            fila += 1
wq[f'A{fila + 1}'] = ('Un polígono no se puede partir entre dos quintiles, por eso cada quintil no tiene exactamente el 20% de la población. '
                      'El mínimo del indicador en cada quintil equivale al corte de ese quintil.')
for colx, wd in zip('ABCD', [8, 22, 36, 22]):
    wq.column_dimensions[colx].width = wd
wq.freeze_panes = 'E2'

# ---- Validación contra el microdato de personas
wv_ = wb.create_sheet('Validación')
wv_.append(['CUT', 'Comuna', 'Población 18+ en polígonos', 'Población 18+ en microdato', 'Cobertura de los polígonos',
            'Escolaridad 18+ ponderada (polígonos)', 'Escolaridad 18+ en microdato', 'Diferencia (años)', 'Chequeo'])
for i, cut in enumerate(CUTS):
    r = 2 + i
    mm = MIC.get(cut)
    esc = rng('esc')
    wv_.append([cut, NOMBRE[cut], f'=SUMIFS({rng("pob18")},{rng("CUT")},$A{r})', None if mm is None else mm['pob18_mic'],
                f'=IF(D{r}="","",C{r}/D{r})',
                f'=SUMPRODUCT(--({rng("CUT")}=$A{r}),{rng("w")},{esc})/SUMPRODUCT(--({rng("CUT")}=$A{r}),--ISNUMBER({esc}),{rng("w")})',
                None if mm is None else round(mm['esc_mic'], 6),
                f'=IF(G{r}="","",F{r}-G{r})',
                f'=IF(G{r}="","Sin dato en el microdato público",IF(ABS(H{r})<={CR["tol"]},"OK","REVISAR"))'])
    wv_[f'E{r}'].number_format = '0.0%'
    for cc in 'FGH':
        wv_[f'{cc}{r}'].number_format = '0.000'
    for cc in 'CD':
        wv_[f'{cc}{r}'].number_format = '#,##0'
nr = 2 + len(CUTS) + 1
wv_[f'A{nr}'] = ('La escolaridad ponderada (promedio de prom_escolaridad18 con peso población 18+) debería reproducir la escolaridad media de las '
                 'personas de 18+ del microdato (excluye el código -99, sin dato). La población del microdato es algo mayor porque incluye a personas '
                 'censadas fuera de un polígono (contenedor comunal: situación de calle y operativos especiales). Las comunas que no tienen dato en '
                 'el microdato público aparecen con otro código de comuna por tamaño (bajo umbral).')
for colx, wd in zip('ABCDEFGHI', [8, 22, 14, 14, 12, 14, 14, 12, 30]):
    wv_.column_dimensions[colx].width = wd

# ---- Diccionario
DICC = [
    ('Columnas grises (insumos)', 'Datos de la base manzana-entidad del Censo 2024 (INE), tal como vienen. Celda vacía = dato suprimido por el INE (menos de 4 casos). Son los únicos valores escritos a mano; todo lo demás es fórmula.'),
    ('Población 18+ (ponderador)', 'n_per − (n_edad_0_5 + n_edad_6_13 + n_edad_14_17). Los tramos de edad no vienen suprimidos en esta base.'),
    ('Peso usado', 'La población 18+ si alcanza el mínimo de la hoja Criterios (4); si no, 0. Con peso 0 el polígono queda sin índice pero sigue en la base.'),
    ('Fila de la comuna en Parámetros', 'Auxiliar: posición de la comuna en la hoja Parámetros, para leer su media y desviación estándar.'),
    ('Componentes (7 columnas verdes)', 'Mismas fórmulas del documento metodológico, redondeadas a un decimal. Vacías si el numerador está suprimido o el denominador es 0.'),
    ('z (7 columnas verdes)', 'signo × (valor − μ de la comuna) / σ de la comuna, con μ y σ ponderados de la hoja Parámetros. Vacío si la comuna no alcanza el n efectivo mínimo en ese componente.'),
    ('Índice 7 componentes', 'Promedio simple de los z disponibles, si hay al menos 4. 0 = adulto promedio de la comuna.'),
    ('Índice 6 componentes (sin escolaridad)', 'Igual, sin el componente de escolaridad promedio. Prueba de robustez.'),
    ('Índice de escolaridad 18+ (años)', 'Escolaridad promedio de las personas de 18+ del polígono, sin estandarizar. Prueba de robustez.'),
    ('Peso en percentil', 'Auxiliar: el peso del polígono si tiene ese indicador; 0 si no.'),
    ('Percentil ponderado en la comuna', '% de la población 18+ de la comuna que vive en polígonos con un valor estrictamente menor del indicador.'),
    ('Quintil ponderado (todos los polígonos)', '1 + parte entera de percentil/20, con máximo 5. Cada quintil reúne cerca del 20% de la población 18+ de la comuna (ver hoja Población por quintil).'),
    ('Quintil ponderado (solo manzanas urbanas)', 'Igual, con el percentil calculado solo sobre la población 18+ de las manzanas urbanas. Vacío en las entidades rurales.'),
]
wd_ = wb.create_sheet('Diccionario')
wd_.append(['Columna', 'Descripción'])
for row in DICC:
    wd_.append(list(row))
wd_.column_dimensions['A'].width = 44
wd_.column_dimensions['B'].width = 120

for w in wb.worksheets[1:]:
    for cell in w[1]:
        cell.fill = PatternFill('solid', fgColor=AZUL)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    w.row_dimensions[1].height = 45
for cell in wp[2]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical='center')
wp.row_dimensions[2].height = 45
for w in [wcr, wp, wco, wq, wv_]:
    for row in w.iter_rows(min_row=w.max_row, max_row=w.max_row):
        for cell in row:
            if isinstance(cell.value, str) and not cell.value.startswith('='):
                cell.font = Font(italic=True, color='5B6670')

os.makedirs(SALIDA_DIR, exist_ok=True)
wb.save(SALIDA)

# ---- que Excel calcule y guarde los valores
ps = (f"$x=New-Object -ComObject Excel.Application; $x.Visible=$false; $x.DisplayAlerts=$false; "
      f"$b=$x.Workbooks.Open('{os.path.abspath(SALIDA)}'); $x.CalculateFull(); $b.Save(); $b.Close($true); $x.Quit(); 'calculado'")
print(subprocess.run(['powershell.exe', '-NoProfile', '-Command', ps], capture_output=True, text=True).stdout.strip())

# ---- validación: valores calculados por Excel contra el cálculo en Python
wv = load_workbook(SALIDA, data_only=True)
m = wv['Manzanas']
CHEQ = ['i7', 'i6', 'iesc'] + [f'{k}_{t}' for k in IDX for t in ['pct', 'q', 'qurb']]
py = {str(int(a)): row for a, row in zip(d['MANZENT'].to_list(), d.select(CHEQ).rows())}
norm = lambda v: None if v in ('', None) else round(float(v), 6)
dif = {k: 0 for k in CHEQ}
errores = 0
for r in range(R0, R1 + 1):
    vals = {k: m[ref(k, r)].value for k in CHEQ}
    if any(isinstance(v, str) and v.startswith('#') for v in vals.values()):
        errores += 1
    p = py[m[ref('MANZENT', r)].value]
    for j, k in enumerate(CHEQ):
        if norm(vals[k]) != norm(p[j]):
            dif[k] += 1
print('celdas con error de Excel:', errores, '| diferencias Excel vs Python:', dif)
hp = wv['Parámetros']
dmax = max(abs(hp[f'{PCOL[(k, "mu")]}{P0 + i}'].value - par[(cut, k)]['mu']) for i, cut in enumerate(CUTS) for k in ORDEN_COMP)
print('máxima diferencia de μ Excel vs Python:', dmax)
# cifras que cita el documento metodológico (índice de 7 componentes, todos los polígonos de la región, mediana ponderada)
v = d.filter(c('i7').is_not_null()).sort('i7')
cw = v['w'].cum_sum() / v['w'].sum()
print('documento: mínimo', v['i7'].min(), '| máximo', v['i7'].max(), '| mediana ponderada', v.filter(cw >= 0.5)['i7'][0],
      '| comunas sin índice por n efectivo:', sorted({NOMBRE[cut] for cut in CUTS if sum(par[(cut, k)]['valido'] for k in ORDEN_COMP) < MIN_COMP}))
hv = wv['Validación']
print('validación escolaridad:', [(hv[f'B{r}'].value, round(hv[f'H{r}'].value, 3) if isinstance(hv[f'H{r}'].value, float) else None, hv[f'I{r}'].value) for r in range(2, 2 + len(CUTS))])
print('Excel:', SALIDA)
